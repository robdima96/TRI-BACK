"""Build an editable PowerPoint reproduction of the Month 1 Prototype Data Flow diagram."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------------------
# Palette
# ---------------------------------------------------------------------------
BG          = RGBColor(0x1C, 0x1E, 0x2E)
BOX_FILL    = RGBColor(0x28, 0x2A, 0x3C)
BOX_BORDER  = RGBColor(0x4A, 0x4D, 0x62)
SECTION_BDR = RGBColor(0x35, 0x38, 0x4C)
CALLOUT_FILL= RGBColor(0x3E, 0x50, 0x62)
CALLOUT_BDR = RGBColor(0x5C, 0x6E, 0x80)
NOTIMPL_FILL= RGBColor(0x22, 0x24, 0x34)
NOTIMPL_BDR = RGBColor(0x3A, 0x3D, 0x50)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY       = RGBColor(0xB0, 0xB2, 0xC4)
DGRAY       = RGBColor(0x70, 0x72, 0x88)
ACCENT      = RGBColor(0x5B, 0x8D, 0xEE)
RED_X       = RGBColor(0xD0, 0x44, 0x44)
ARROW_CLR   = RGBColor(0x8A, 0x8C, 0x9E)

# ---------------------------------------------------------------------------
# Presentation setup
# ---------------------------------------------------------------------------
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

bg_fill = slide.background.fill
bg_fill.solid()
bg_fill.fore_color.rgb = BG

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _rr(left, top, w, h, fill_c, bdr_c, bdr_w=Pt(1)):
    """Add a rounded rectangle."""
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill_c
    if bdr_c:
        s.line.color.rgb = bdr_c
        s.line.width = bdr_w
    else:
        s.line.fill.background()
    # smaller corner radius
    s.adjustments[0] = 0.06
    tf = s.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    return s

def _tb(left, top, w, h):
    """Add a text box and return its text_frame."""
    s = slide.shapes.add_textbox(left, top, w, h)
    s.text_frame.word_wrap = True
    return s.text_frame

def _run(para, text, sz, bold=False, italic=False, color=WHITE):
    r = para.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return r

def _para(tf, text, sz, bold=False, italic=False, color=WHITE, align=PP_ALIGN.CENTER):
    p = tf.add_paragraph()
    p.alignment = align
    _run(p, text, sz, bold, italic, color)
    p.space_before = Pt(0)
    p.space_after  = Pt(0)
    return p

def _first(tf, text, sz, bold=False, italic=False, color=WHITE, align=PP_ALIGN.CENTER):
    """Set text on the first (already-existing) paragraph."""
    p = tf.paragraphs[0]
    p.alignment = align
    _run(p, text, sz, bold, italic, color)
    p.space_before = Pt(0)
    p.space_after  = Pt(0)
    return p

def _arrow_down(cx, y1, y2):
    """Vertical downward arrow (thin line + triangle head)."""
    line_w = Inches(0.018)
    head_w = Inches(0.14)
    head_h = Inches(0.10)
    line_h = (y2 - y1) - head_h
    if line_h < 0:
        line_h = Emu(0)
    # stem
    stem = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx - line_w//2, y1, line_w, line_h)
    stem.fill.solid(); stem.fill.fore_color.rgb = ARROW_CLR
    stem.line.fill.background()
    # head
    tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                  cx - head_w//2, y1 + line_h, head_w, head_h)
    tri.fill.solid(); tri.fill.fore_color.rgb = ARROW_CLR
    tri.line.fill.background()
    tri.rotation = 180.0

def _arrow_right(x1, cy, x2):
    """Horizontal rightward arrow (thin line + triangle head)."""
    line_h = Inches(0.018)
    head_w = Inches(0.10)
    head_h = Inches(0.14)
    line_w = (x2 - x1) - head_w
    if line_w < 0:
        line_w = Emu(0)
    stem = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x1, cy - line_h//2, line_w, line_h)
    stem.fill.solid(); stem.fill.fore_color.rgb = ARROW_CLR
    stem.line.fill.background()
    tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                  x1 + line_w, cy - head_h//2, head_w, head_h)
    tri.fill.solid(); tri.fill.fore_color.rgb = ARROW_CLR
    tri.line.fill.background()
    tri.rotation = 90.0

def _dashed_rect(left, top, w, h, color):
    """Section boundary – no fill, dashed border."""
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    s.fill.background()
    s.line.color.rgb = color
    s.line.width = Pt(0.75)
    s.line.dash_style = 4  # dash
    s.adjustments[0] = 0.02
    return s

# ---------------------------------------------------------------------------
# TITLE
# ---------------------------------------------------------------------------
tf = _tb(Inches(0.5), Inches(0.12), Inches(12.3), Inches(0.55))
_first(tf, "DigiMSK Chatbot \u2014 Month 1 Prototype Data Flow", 26, bold=True)

tf = _tb(Inches(0.5), Inches(0.60), Inches(12.3), Inches(0.32))
_first(tf, "Linear single-turn pipeline with service stubs", 13, color=LGRAY)

# ---------------------------------------------------------------------------
# USER BOX
# ---------------------------------------------------------------------------
user_l, user_t, user_w, user_h = Inches(5.85), Inches(1.05), Inches(1.5), Inches(0.45)
user = _rr(user_l, user_t, user_w, user_h, BOX_FILL, BOX_BORDER)
tf = user.text_frame
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
_run(tf.paragraphs[0], "\U0001F464  User", 11, bold=True)

user_cx = user_l + user_w // 2

# ChatRequest / ChatResponse labels
tf = _tb(Inches(2.6), Inches(1.15), Inches(3.1), Inches(0.28))
_first(tf, "ChatRequest (session_id + message)", 8.5, color=LGRAY, align=PP_ALIGN.RIGHT)

tf = _tb(Inches(7.5), Inches(1.15), Inches(4.0), Inches(0.28))
_first(tf, "ChatResponse (response + citations + escalated + safety_reason)", 8.5, color=LGRAY, align=PP_ALIGN.LEFT)

# Arrow: User -> FastAPI
_arrow_down(user_cx, user_t + user_h, Inches(1.72))

# ---------------------------------------------------------------------------
# API LAYER
# ---------------------------------------------------------------------------
api_sect = _dashed_rect(Inches(1.5), Inches(1.72), Inches(10.3), Inches(0.62), SECTION_BDR)

tf = _tb(Inches(0.3), Inches(1.74), Inches(1.5), Inches(0.28))
_first(tf, "API LAYER", 10, bold=True, color=LGRAY, align=PP_ALIGN.LEFT)

fa_l, fa_t, fa_w, fa_h = Inches(5.0), Inches(1.76), Inches(3.2), Inches(0.52)
fa = _rr(fa_l, fa_t, fa_w, fa_h, BOX_FILL, BOX_BORDER)
tf = fa.text_frame
_first(tf, "FastAPI Endpoint", 11, bold=True)
_para(tf, "POST /api/v1/chat", 9, color=LGRAY)

fa_cx = fa_l + fa_w // 2

# Arrow: FastAPI -> orchestrator
_arrow_down(fa_cx, fa_t + fa_h, Inches(2.55))

# ---------------------------------------------------------------------------
# ORCHESTRATOR SECTION (boundary)
# ---------------------------------------------------------------------------
orch_t = Inches(2.55)
orch_b = Inches(6.10)
_dashed_rect(Inches(1.5), orch_t, Inches(10.3), orch_b - orch_t, SECTION_BDR)

tf = _tb(Inches(0.3), Inches(2.55), Inches(2.7), Inches(0.26))
_first(tf, "ORCHESTRATOR", 10, bold=True, color=LGRAY, align=PP_ALIGN.LEFT)
tf = _tb(Inches(0.3), Inches(2.78), Inches(2.8), Inches(0.25))
_first(tf, "LangGraph Linear Graph \u2014 Single Turn Only", 7.5, color=DGRAY, align=PP_ALIGN.LEFT)

# ---------------------------------------------------------------------------
# PIPELINE STEPS
# ---------------------------------------------------------------------------
PX = Inches(3.8)
PW = Inches(5.0)
Y0 = 2.70
DY = 0.72

steps = [
    ("ingest_input",
     "Keyword risk scan (policy.py)",
     "Matches against: suicide, self-harm, chest pain, numbness, loss of bladder control"),
    ("retrieve_evidence",
     "RAG Service (STUB)",
     "Returns hardcoded evidence. Future: vector DB retrieval"),
    ("generate_draft",
     "Generator Service (STUB)",
     "Returns canned response. Future: Llama-3-8B"),
    ("policy_gate",
     "Safety Override",
     "If risk flags present \u2192 override draft with escalation message"),
    ("finalize_response",
     "Package output",
     None),
]

box_tops = []
box_heights = []

for i, (name, sub, desc) in enumerate(steps):
    y = Inches(Y0 + i * DY)
    h = Inches(0.60) if desc else Inches(0.48)
    box = _rr(PX, y, PW, h, BOX_FILL, BOX_BORDER)
    tf = box.text_frame
    tf.margin_top = Pt(3)
    tf.margin_bottom = Pt(3)
    _first(tf, name, 11, bold=True)
    _para(tf, sub, 8.5, italic=True, color=LGRAY)
    if desc:
        _para(tf, desc, 7.5, color=DGRAY)
    box_tops.append(y)
    box_heights.append(h)

pipe_cx = PX + PW // 2

for i in range(len(steps) - 1):
    _arrow_down(pipe_cx, box_tops[i] + box_heights[i], box_tops[i + 1])

# ---------------------------------------------------------------------------
# SIDE CALLOUT BOXES  (RAG + GENERATOR)
# ---------------------------------------------------------------------------
CO_L = Inches(9.3)
CO_W = Inches(2.6)
CO_H = Inches(0.55)

# RAG callout – aligned with retrieve_evidence (index 1)
rag_t = box_tops[1] + (box_heights[1] - CO_H) // 2
rag = _rr(CO_L, rag_t, CO_W, CO_H, CALLOUT_FILL, CALLOUT_BDR)
tf = rag.text_frame
tf.margin_top = Pt(2)
_first(tf, "RAG (STUB)", 10, bold=True, color=ACCENT)
_para(tf, "Hardcoded evidence \u2014 no real\nretrieval yet", 7.5, color=LGRAY)

_arrow_right(PX + PW, box_tops[1] + box_heights[1] // 2, CO_L)

# GENERATOR callout – aligned with generate_draft (index 2)
gen_t = box_tops[2] + (box_heights[2] - CO_H) // 2
gen = _rr(CO_L, gen_t, CO_W, CO_H, CALLOUT_FILL, CALLOUT_BDR)
tf = gen.text_frame
tf.margin_top = Pt(2)
_first(tf, "GENERATOR (STUB)", 10, bold=True, color=ACCENT)
_para(tf, "Canned response \u2014 no LLM yet", 7.5, color=LGRAY)

_arrow_right(PX + PW, box_tops[2] + box_heights[2] // 2, CO_L)

# ---------------------------------------------------------------------------
# RESPONSE SECTION
# ---------------------------------------------------------------------------
resp_t = Inches(6.20)
_dashed_rect(Inches(0.4), resp_t, Inches(12.5), Inches(0.95), SECTION_BDR)

tf = _tb(Inches(0.3), Inches(6.15), Inches(1.5), Inches(0.26))
_first(tf, "RESPONSE", 10, bold=True, color=LGRAY, align=PP_ALIGN.LEFT)

# ChatResponse label above response boxes
tf = _tb(Inches(4.2), Inches(6.18), Inches(4.8), Inches(0.22))
_first(tf, "ChatResponse (response + citations + escalated + safety_reason)", 7.5, color=LGRAY)

# Arrow down from finalize_response to response section
last_bot = box_tops[-1] + box_heights[-1]
_arrow_down(pipe_cx, last_bot, resp_t + Inches(0.08))

# Three "not implemented" boxes
ni_y = Inches(6.50)
ni_h = Inches(0.48)

def _not_impl(left, width, label, has_x=True):
    box = _rr(left, ni_y, width, ni_h, NOTIMPL_FILL, NOTIMPL_BDR)
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    if has_x:
        _run(p, "\u2715 ", 11, bold=True, color=RED_X)
    _run(p, label, 9, color=LGRAY)
    return box

_not_impl(Inches(0.6),  Inches(3.2), "ENCODER \u2014 Not implemented")
_not_impl(Inches(4.1),  Inches(3.6), "SYNTHESIZER \u2014 Not implemented")
_not_impl(Inches(8.0),  Inches(4.7), "Conversation History \u2014 Not implemented\n(single-turn only)", has_x=False)

# ---------------------------------------------------------------------------
# FOOTER
# ---------------------------------------------------------------------------
tf = _tb(Inches(0.5), Inches(7.18), Inches(12.3), Inches(0.25))
_first(tf, "Ardern Lab, UBC \u2014 April 2026", 9, color=DGRAY)

# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
out = r"C:\ROBS STUFF\UBC Postdoctoral Fellowship\DigiMSKbot\bot\docs\month1\Month1_Prototype_DataFlow.pptx"
prs.save(out)
print(f"Saved to {out}")
