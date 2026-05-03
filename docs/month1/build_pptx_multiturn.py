"""Build editable PowerPoint for the Multi-Turn Conversational Data Flow."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------------------
# Palette
# ---------------------------------------------------------------------------
BG          = RGBColor(0x1C, 0x1E, 0x2E)
BOX_FILL    = RGBColor(0x28, 0x2A, 0x3C)
BOX_BORDER  = RGBColor(0x4A, 0x4D, 0x62)
SECTION_BDR = RGBColor(0x35, 0x38, 0x4C)
CALLOUT_FILL= RGBColor(0x3E, 0x50, 0x62)
CALLOUT_BDR = RGBColor(0x5C, 0x6E, 0x80)
NOTIMPL_FILL= RGBColor(0x22, 0x24, 0x34)
NOTIMPL_BDR = RGBColor(0x3A, 0x3D, 0x50)
IMPL_FILL   = RGBColor(0x1E, 0x30, 0x28)
IMPL_BDR    = RGBColor(0x3A, 0x5D, 0x4A)
STORE_FILL  = RGBColor(0x2A, 0x3A, 0x4E)
STORE_BDR   = RGBColor(0x4E, 0x6A, 0x80)

WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY       = RGBColor(0xB0, 0xB2, 0xC4)
DGRAY       = RGBColor(0x70, 0x72, 0x88)
ACCENT      = RGBColor(0x5B, 0x8D, 0xEE)
RED_X       = RGBColor(0xD0, 0x44, 0x44)
GREEN_CHK   = RGBColor(0x40, 0xB0, 0x40)
ARROW_CLR   = RGBColor(0x8A, 0x8C, 0x9E)
LOOP_CLR    = RGBColor(0x5B, 0x8D, 0xEE)

# ---------------------------------------------------------------------------
# Presentation setup
# ---------------------------------------------------------------------------
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

bg_fill = slide.background.fill
bg_fill.solid()
bg_fill.fore_color.rgb = BG

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _rr(left, top, w, h, fill_c, bdr_c, bdr_w=Pt(1)):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill_c
    if bdr_c:
        s.line.color.rgb = bdr_c
        s.line.width = bdr_w
    else:
        s.line.fill.background()
    s.adjustments[0] = 0.06
    tf = s.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    return s

def _tb(left, top, w, h):
    s = slide.shapes.add_textbox(left, top, w, h)
    s.text_frame.word_wrap = True
    return s.text_frame

def _run(para, text, sz, bold=False, italic=False, color=WHITE):
    r = para.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return r

def _para(tf, text, sz, bold=False, italic=False, color=WHITE, align=PP_ALIGN.CENTER):
    p = tf.add_paragraph()
    p.alignment = align
    _run(p, text, sz, bold, italic, color)
    p.space_before = Pt(0)
    p.space_after  = Pt(0)
    return p

def _first(tf, text, sz, bold=False, italic=False, color=WHITE, align=PP_ALIGN.CENTER):
    p = tf.paragraphs[0]
    p.alignment = align
    _run(p, text, sz, bold, italic, color)
    p.space_before = Pt(0)
    p.space_after  = Pt(0)
    return p

def _arrow_down(cx, y1, y2, hw=Inches(0.12), hh=Inches(0.08)):
    line_w = Inches(0.018)
    line_h = (y2 - y1) - hh
    if line_h < 0:
        line_h = Emu(0)
    stem = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                   cx - line_w // 2, y1, line_w, line_h)
    stem.fill.solid(); stem.fill.fore_color.rgb = ARROW_CLR
    stem.line.fill.background()
    tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                  cx - hw // 2, y1 + line_h, hw, hh)
    tri.fill.solid(); tri.fill.fore_color.rgb = ARROW_CLR
    tri.line.fill.background()
    tri.rotation = 180.0

def _arrow_right(x1, cy, x2, clr=ARROW_CLR):
    line_h = Inches(0.018)
    head_w = Inches(0.10)
    head_h = Inches(0.14)
    line_w = (x2 - x1) - head_w
    if line_w < 0:
        line_w = Emu(0)
    stem = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                   x1, cy - line_h // 2, line_w, line_h)
    stem.fill.solid(); stem.fill.fore_color.rgb = clr
    stem.line.fill.background()
    tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                  x1 + line_w, cy - head_h // 2, head_w, head_h)
    tri.fill.solid(); tri.fill.fore_color.rgb = clr
    tri.line.fill.background()
    tri.rotation = 90.0

def _arrow_left(x1, cy, x2, clr=ARROW_CLR):
    """Arrow pointing LEFT from x1 to x2  (x2 < x1)."""
    line_h = Inches(0.018)
    head_w = Inches(0.10)
    head_h = Inches(0.14)
    total  = x1 - x2
    line_w = total - head_w
    if line_w < 0:
        line_w = Emu(0)
    stem = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                   x2 + head_w, cy - line_h // 2, line_w, line_h)
    stem.fill.solid(); stem.fill.fore_color.rgb = clr
    stem.line.fill.background()
    tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                  x2, cy - head_h // 2, head_w, head_h)
    tri.fill.solid(); tri.fill.fore_color.rgb = clr
    tri.line.fill.background()
    tri.rotation = 270.0

def _arrow_up(cx, y_bottom, y_top, clr=LOOP_CLR):
    line_w = Inches(0.018)
    hw = Inches(0.12)
    hh = Inches(0.08)
    line_h = (y_bottom - y_top) - hh
    if line_h < 0:
        line_h = Emu(0)
    stem = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                   cx - line_w // 2, y_top + hh, line_w, line_h)
    stem.fill.solid(); stem.fill.fore_color.rgb = clr
    stem.line.fill.background()
    tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                  cx - hw // 2, y_top, hw, hh)
    tri.fill.solid(); tri.fill.fore_color.rgb = clr
    tri.line.fill.background()

def _dashed_rect(left, top, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    s.fill.background()
    s.line.color.rgb = color
    s.line.width = Pt(0.75)
    s.line.dash_style = 4
    s.adjustments[0] = 0.02
    return s

# ---------------------------------------------------------------------------
# TITLE
# ---------------------------------------------------------------------------
tf = _tb(Inches(0.5), Inches(0.12), Inches(12.3), Inches(0.55))
_first(tf, "DigiMSK Chatbot \u2014 Multi-Turn Conversational Data Flow", 26, bold=True)

tf = _tb(Inches(0.5), Inches(0.60), Inches(12.3), Inches(0.32))
_first(tf, "Looping pipeline with session memory and live services", 13, color=LGRAY)

# ---------------------------------------------------------------------------
# USER BOX
# ---------------------------------------------------------------------------
user_l, user_t, user_w, user_h = Inches(5.85), Inches(1.05), Inches(1.5), Inches(0.45)
user = _rr(user_l, user_t, user_w, user_h, BOX_FILL, BOX_BORDER)
tf = user.text_frame
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
_run(tf.paragraphs[0], "\U0001F464  User", 11, bold=True)

user_cx = user_l + user_w // 2

tf = _tb(Inches(2.6), Inches(1.15), Inches(3.1), Inches(0.28))
_first(tf, "ChatRequest (session_id + message)", 8.5, color=LGRAY, align=PP_ALIGN.RIGHT)

tf = _tb(Inches(7.5), Inches(1.15), Inches(4.0), Inches(0.28))
_first(tf, "ChatResponse (response + citations + escalated + safety_reason)",
       8.5, color=LGRAY, align=PP_ALIGN.LEFT)

_arrow_down(user_cx, user_t + user_h, Inches(1.72))

# ---------------------------------------------------------------------------
# API LAYER
# ---------------------------------------------------------------------------
_dashed_rect(Inches(1.5), Inches(1.72), Inches(10.3), Inches(0.62), SECTION_BDR)

tf = _tb(Inches(0.3), Inches(1.74), Inches(1.5), Inches(0.28))
_first(tf, "API LAYER", 10, bold=True, color=LGRAY, align=PP_ALIGN.LEFT)

fa_l, fa_t, fa_w, fa_h = Inches(5.0), Inches(1.76), Inches(3.2), Inches(0.52)
fa = _rr(fa_l, fa_t, fa_w, fa_h, BOX_FILL, BOX_BORDER)
tf = fa.text_frame
_first(tf, "FastAPI Endpoint", 11, bold=True)
_para(tf, "POST /api/v1/chat", 9, color=LGRAY)

fa_cx = fa_l + fa_w // 2
_arrow_down(fa_cx, fa_t + fa_h, Inches(2.42))

# ---------------------------------------------------------------------------
# ORCHESTRATOR SECTION
# ---------------------------------------------------------------------------
orch_t = Inches(2.42)
orch_b = Inches(6.12)
_dashed_rect(Inches(1.2), orch_t, Inches(10.8), orch_b - orch_t, SECTION_BDR)

tf = _tb(Inches(0.3), Inches(2.42), Inches(2.7), Inches(0.26))
_first(tf, "ORCHESTRATOR", 10, bold=True, color=LGRAY, align=PP_ALIGN.LEFT)
tf = _tb(Inches(0.3), Inches(2.64), Inches(2.8), Inches(0.25))
_first(tf, "LangGraph Cyclic Graph \u2014 Multi-Turn", 7.5, color=DGRAY, align=PP_ALIGN.LEFT)

# ---------------------------------------------------------------------------
# PIPELINE STEPS (7)
# ---------------------------------------------------------------------------
PX = Inches(4.2)
PW = Inches(4.6)
Y0 = 2.55
DY = 0.50

steps = [
    ("load_history",
     "Session Store Read (history.py)",
     "Retrieve prior turns + state from session-keyed store"),
    ("ingest_input",
     "Keyword risk scan (policy.py)",
     "Risk keywords evaluated with conversation context"),
    ("retrieve_evidence",
     "RAG Service",
     "Vector DB retrieval with query reformulation"),
    ("generate_draft",
     "Generator Service (Llama-3-8B)",
     "Context-aware response using conversation history"),
    ("policy_gate",
     "Safety Override",
     "Risk flags from current + prior turns \u2192 escalation"),
    ("finalize_response",
     "Package output",
     None),
    ("save_turn",
     "Session Store Write (history.py)",
     "Persist turn + updated state to session store"),
]

box_tops = []
box_heights = []

for i, (name, sub, desc) in enumerate(steps):
    y = Inches(Y0 + i * DY)
    h = Inches(0.40) if desc else Inches(0.30)
    box = _rr(PX, y, PW, h, BOX_FILL, BOX_BORDER)
    tf = box.text_frame
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    _first(tf, name, 10, bold=True)
    _para(tf, sub, 8, italic=True, color=LGRAY)
    if desc:
        _para(tf, desc, 7, color=DGRAY)
    box_tops.append(y)
    box_heights.append(h)

pipe_cx = PX + PW // 2

for i in range(len(steps) - 1):
    _arrow_down(pipe_cx,
                box_tops[i] + box_heights[i],
                box_tops[i + 1],
                hw=Inches(0.10), hh=Inches(0.06))

# ---------------------------------------------------------------------------
# SESSION STORE  (tall left-side element)
# ---------------------------------------------------------------------------
STORE_L = Inches(1.5)
STORE_W = Inches(2.2)
store_top = box_tops[0]
store_bot = box_tops[6] + box_heights[6]
STORE_H = store_bot - store_top

store = _rr(STORE_L, store_top, STORE_W, STORE_H, STORE_FILL, STORE_BDR, Pt(1.5))
store.adjustments[0] = 0.03
tf = store.text_frame
tf.word_wrap = True
tf.auto_size = None
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
_first(tf, "SESSION STORE", 12, bold=True, color=ACCENT)
_para(tf, "Redis / in-memory dict", 8, color=LGRAY)
_para(tf, "Session-keyed", 7.5, color=DGRAY)
_para(tf, "conversation state", 7.5, color=DGRAY)
_para(tf, "", 6)
_para(tf, "\u21BB  next turn", 10, bold=True, color=LOOP_CLR)

store_right = STORE_L + STORE_W

# Arrow: store → load_history  (read)
load_cy = box_tops[0] + box_heights[0] // 2
_arrow_right(store_right, load_cy, PX, clr=LOOP_CLR)

# Arrow: save_turn → store  (write)
save_cy = box_tops[6] + box_heights[6] // 2
_arrow_left(PX, save_cy, store_right, clr=LOOP_CLR)

# "read" / "write" labels near the arrows
tf = _tb(store_right + Inches(0.02), load_cy - Inches(0.24), Inches(0.40), Inches(0.20))
_first(tf, "read", 7, italic=True, color=DGRAY)

tf = _tb(store_right + Inches(0.02), save_cy + Inches(0.04), Inches(0.40), Inches(0.20))
_first(tf, "write", 7, italic=True, color=DGRAY)

# Loop-back indicator: upward arrow on far left of store
loop_x = STORE_L - Inches(0.25)
_arrow_up(loop_x, save_cy, load_cy, clr=LOOP_CLR)

tf = _tb(Inches(0.08), Inches((2.55 + 5.55) / 2 - 0.10), Inches(1.1), Inches(0.40))
_first(tf, "\u21BB  loop", 8, bold=True, color=LOOP_CLR, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------------------
# RIGHT-SIDE CALLOUTS  (RAG + GENERATOR — no longer stubs)
# ---------------------------------------------------------------------------
CO_L = Inches(9.3)
CO_W = Inches(2.8)
CO_H = Inches(0.52)

# RAG — aligned with retrieve_evidence (step 2)
rag_cy = box_tops[2] + box_heights[2] // 2
rag_t = box_tops[2] + (box_heights[2] - CO_H) // 2
rag = _rr(CO_L, rag_t, CO_W, CO_H, CALLOUT_FILL, CALLOUT_BDR)
tf = rag.text_frame
tf.margin_top = Pt(2)
_first(tf, "RAG", 10, bold=True, color=ACCENT)
_para(tf, "Vector DB retrieval \u2014 FAISS / pgvector", 7.5, color=LGRAY)

_arrow_right(PX + PW, rag_cy, CO_L)

# GENERATOR — aligned with generate_draft (step 3)
gen_cy = box_tops[3] + box_heights[3] // 2
gen_t = box_tops[3] + (box_heights[3] - CO_H) // 2
gen = _rr(CO_L, gen_t, CO_W, CO_H, CALLOUT_FILL, CALLOUT_BDR)
tf = gen.text_frame
tf.margin_top = Pt(2)
_first(tf, "GENERATOR", 10, bold=True, color=ACCENT)
_para(tf, "Llama-3-8B-Instruct with DPO fine-tuning", 7.5, color=LGRAY)

_arrow_right(PX + PW, gen_cy, CO_L)

# ---------------------------------------------------------------------------
# RESPONSE SECTION
# ---------------------------------------------------------------------------
resp_t = Inches(6.22)
_dashed_rect(Inches(0.4), resp_t, Inches(12.5), Inches(0.88), SECTION_BDR)

tf = _tb(Inches(0.3), Inches(6.18), Inches(1.5), Inches(0.26))
_first(tf, "RESPONSE", 10, bold=True, color=LGRAY, align=PP_ALIGN.LEFT)

tf = _tb(Inches(4.2), Inches(6.20), Inches(4.8), Inches(0.22))
_first(tf, "ChatResponse (response + citations + escalated + safety_reason)",
       7.5, color=LGRAY)

last_bot = box_tops[-1] + box_heights[-1]
_arrow_down(pipe_cx, last_bot, resp_t + Inches(0.08))

ni_y = Inches(6.50)
ni_h = Inches(0.45)

def _status_box(left, width, label, implemented=False):
    fill = IMPL_FILL if implemented else NOTIMPL_FILL
    bdr  = IMPL_BDR  if implemented else NOTIMPL_BDR
    box = _rr(left, ni_y, width, ni_h, fill, bdr)
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    if implemented:
        _run(p, "\u2713 ", 11, bold=True, color=GREEN_CHK)
    else:
        _run(p, "\u2715 ", 11, bold=True, color=RED_X)
    _run(p, label, 9, color=LGRAY)
    return box

_status_box(Inches(0.6),  Inches(3.2),
            "ENCODER \u2014 Not implemented", implemented=False)
_status_box(Inches(4.1),  Inches(3.6),
            "SYNTHESIZER \u2014 Not implemented", implemented=False)
_status_box(Inches(8.0),  Inches(4.7),
            "Conversation History \u2014 Implemented", implemented=True)

# ---------------------------------------------------------------------------
# FOOTER
# ---------------------------------------------------------------------------
tf = _tb(Inches(0.5), Inches(7.18), Inches(12.3), Inches(0.25))
_first(tf, "Ardern Lab, UBC \u2014 April 2026", 9, color=DGRAY)

# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
out = r"C:\ROBS STUFF\UBC Postdoctoral Fellowship\DigiMSKbot\bot\docs\month1\MultiTurn_DataFlow.pptx"
prs.save(out)
print(f"Saved to {out}")
